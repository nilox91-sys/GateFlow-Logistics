import io
from datetime import datetime, timezone, date
from typing import List
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import pandas as pd

from ..database import get_db
from ..models import Transito, StatoTransito
from ..schemas import CheckInRequest, UpdateStatusRequest, UpdateTransitoRequest, TransitoResponse
from .websockets import manager
from .auth import get_current_user
from ..utils.audit import add_audit_log

router = APIRouter()

def serialize_transito(t: Transito) -> dict:
    return {
        "id": t.id,
        "targa": t.targa,
        "vettore": t.vettore,
        "codice_linea": t.codice_linea,
        "codice_vred_tme": t.codice_vred_tme,
        "cliente": t.cliente,
        "partenza": t.partenza,
        "autista": t.autista,
        "telefono": t.telefono,
        "targa_smr": t.targa_smr,
        "tipo_operazione": t.tipo_operazione.value if t.tipo_operazione else None,
        "ddt": t.ddt,
        "colli": t.colli,
        "peso": t.peso,
        "bi": t.bi,
        "ci": t.ci,
        "sacco": t.sacco,
        "percentuale_sfu": t.percentuale_sfu,
        "pe": t.pe,
        "rr": t.rr,
        "me": t.me,
        "sigillo1": t.sigillo1,
        "sigillo2": t.sigillo2,
        "note": t.note,
        "timestamp_in": t.timestamp_in.isoformat() if t.timestamp_in else None,
        "timestamp_out": t.timestamp_out.isoformat() if t.timestamp_out else None,
        "molo": t.molo,
        "stato": t.stato.value if t.stato else None,
    }

@router.post("/check-in", response_model=TransitoResponse, tags=["Portineria"])
async def check_in(payload: CheckInRequest, db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    nuovo = Transito(
        targa=payload.targa.upper().strip(),
        vettore=payload.vettore.strip(),
        codice_linea=payload.codice_linea,
        codice_vred_tme=payload.codice_vred_tme,
        cliente=payload.cliente,
        partenza=payload.partenza,
        autista=payload.autista,
        telefono=payload.telefono,
        targa_smr=payload.targa_smr,
        tipo_operazione=payload.tipo_operazione,
        ddt=payload.ddt,
        colli=payload.colli,
        peso=payload.peso,
        bi=payload.bi,
        ci=payload.ci,
        sacco=payload.sacco,
        percentuale_sfu=payload.percentuale_sfu,
        pe=payload.pe,
        rr=payload.rr,
        me=payload.me,
        sigillo1=payload.sigillo1,
        sigillo2=payload.sigillo2,
        note=payload.note,
        molo=payload.molo,
        timestamp_in=datetime.now(),
        stato=StatoTransito.INGRESSO,
    )
    db.add(nuovo)
    db.commit()
    db.refresh(nuovo)

    add_audit_log(db, current_user, "Check-in", f"Ingresso targa {nuovo.targa} (Vettore: {nuovo.vettore})")
    await manager.broadcast({"event": "check_in", "data": serialize_transito(nuovo)})
    return nuovo

@router.patch("/update-status/{transito_id}", response_model=TransitoResponse, tags=["Magazzino"])
async def update_status(
    transito_id: str,
    payload: UpdateStatusRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    t = db.query(Transito).filter(Transito.id == transito_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Transito non trovato.")

    # Guard: USCITO can only be set if current state is COMPLETATO
    if payload.stato == StatoTransito.USCITO and t.stato != StatoTransito.COMPLETATO:
        raise HTTPException(
            status_code=400,
            detail=f"Impossibile impostare USCITO: lo stato corrente è '{t.stato.value}', deve essere 'COMPLETATO'.",
        )
        
    # Guard: Validation to prevent INGRESSO -> USCITO without passing through intermediate states
    if t.stato == StatoTransito.INGRESSO and payload.stato == StatoTransito.USCITO:
        raise HTTPException(
            status_code=400,
            detail="Transizione non valida: non è possibile passare direttamente da INGRESSO a USCITO."
        )

    t.stato = payload.stato
    update_data = payload.model_dump(exclude_unset=True)
    if 'molo' in update_data:
        t.molo = update_data['molo']
    db.commit()
    db.refresh(t)

    add_audit_log(db, current_user, "Cambio Stato", f"Mezzo {t.targa} -> {t.stato.value} (Molo: {t.molo})")
    await manager.broadcast({"event": "status_update", "data": serialize_transito(t)})
    return t

@router.patch("/transito/{transito_id}", response_model=TransitoResponse, tags=["Gestione"])
async def update_transito(
    transito_id: str,
    payload: UpdateTransitoRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    t = db.query(Transito).filter(Transito.id == transito_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Transito non trovato.")

    update_data = payload.model_dump(exclude_unset=True)
    if update_data.get('targa'):
        t.targa = update_data['targa'].upper().strip()
    if update_data.get('vettore'):
        t.vettore = update_data['vettore'].strip()
    if 'codice_linea' in update_data:
        t.codice_linea = update_data['codice_linea']
    if 'codice_vred_tme' in update_data:
        t.codice_vred_tme = update_data['codice_vred_tme']
    if 'cliente' in update_data:
        t.cliente = update_data['cliente']
    if 'partenza' in update_data:
        t.partenza = update_data['partenza']
    if 'autista' in update_data:
        t.autista = update_data['autista']
    if 'telefono' in update_data:
        t.telefono = update_data['telefono']
    if 'targa_smr' in update_data:
        t.targa_smr = update_data['targa_smr']
    if 'tipo_operazione' in update_data and update_data['tipo_operazione'] is not None:
        t.tipo_operazione = update_data['tipo_operazione']
    if 'molo' in update_data:
        t.molo = update_data['molo']
    if 'stato' in update_data and update_data['stato'] is not None:
        t.stato = update_data['stato']
    if 'ddt' in update_data:
        t.ddt = update_data['ddt']
    if 'colli' in update_data:
        t.colli = update_data['colli']
    if 'peso' in update_data:
        t.peso = update_data['peso']
    if 'bi' in update_data:
        t.bi = update_data['bi']
    if 'ci' in update_data:
        t.ci = update_data['ci']
    if 'sacco' in update_data:
        t.sacco = update_data['sacco']
    if 'percentuale_sfu' in update_data:
        t.percentuale_sfu = update_data['percentuale_sfu']
    if 'pe' in update_data:
        t.pe = update_data['pe']
    if 'rr' in update_data:
        t.rr = update_data['rr']
    if 'me' in update_data:
        t.me = update_data['me']
    if 'sigillo1' in update_data:
        t.sigillo1 = update_data['sigillo1']
    if 'sigillo2' in update_data:
        t.sigillo2 = update_data['sigillo2']
    if 'note' in update_data:
        t.note = update_data['note']

    db.commit()
    db.refresh(t)

    await manager.broadcast({"event": "status_update", "data": serialize_transito(t)})
    return t

@router.post("/check-out/{transito_id}", response_model=TransitoResponse, tags=["Portineria"])
async def check_out(transito_id: str, db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    t = db.query(Transito).filter(Transito.id == transito_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Transito non trovato.")

    if t.stato != StatoTransito.COMPLETATO:
        raise HTTPException(
            status_code=400,
            detail=f"Impossibile effettuare il check-out: stato attuale '{t.stato.value}'. Richiesto: 'COMPLETATO'.",
        )

    t.timestamp_out = datetime.now(tz=timezone.utc)
    t.stato = StatoTransito.USCITO
    db.commit()
    db.refresh(t)

    await manager.broadcast({"event": "check_out", "data": serialize_transito(t)})
    return t

@router.delete("/transito/{transito_id}", tags=["Gestione"])
async def delete_transit(transito_id: str, db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in ["admin", "responsabile"]:
        raise HTTPException(status_code=403, detail="Permesso negato: solo amministratori e responsabili possono eliminare transiti.")
    
    t = db.query(Transito).filter(Transito.id == transito_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Transito non trovato.")
    
    add_audit_log(db, current_user, "Eliminazione", f"Cancellato transito targa {t.targa}")
    db.delete(t)
    db.commit()
    
    await manager.broadcast({"event": "delete", "data": {"id": transito_id}})
    return {"message": "Transito eliminato con successo"}

@router.get("/transits", response_model=List[TransitoResponse], tags=["Dashboard"])
def get_active_transits(db: Session = Depends(get_db)):
    return (
        db.query(Transito)
        .filter(Transito.stato != StatoTransito.USCITO)
        .order_by(Transito.timestamp_in.desc())
        .all()
    )

@router.get("/transits/history", response_model=List[TransitoResponse], tags=["Dashboard"])
def get_history_transits(day: str = None, db: Session = Depends(get_db)):
    """
    Ritorna tutti i transiti di una precisa giornata (inclusi gli USCITO).
    Se `day` non è fornito, usa la giornata odierna. Format: YYYY-MM-DD.
    """
    if day:
        target_date = datetime.strptime(day, "%Y-%m-%d").date()
    else:
        target_date = date.today()
        
    start_dt = datetime.combine(target_date, datetime.min.time(), tzinfo=timezone.utc)
    end_dt = datetime.combine(target_date, datetime.max.time(), tzinfo=timezone.utc)

    return (
        db.query(Transito)
        .filter(Transito.timestamp_in >= start_dt)
        .filter(Transito.timestamp_in <= end_dt)
        .order_by(Transito.timestamp_in.desc())
        .all()
    )

@router.get("/export/daily", tags=["Amministrazione"])
def export_daily(day: str = None, db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    if not current_user.get("can_export"):
        raise HTTPException(status_code=403, detail="Non hai il permesso di esportare i dati.")
        
    if day:
        target_date = datetime.strptime(day, "%Y-%m-%d").date()
    else:
        target_date = date.today()
        
    start_dt = datetime.combine(target_date, datetime.min.time(), tzinfo=timezone.utc)
    end_dt = datetime.combine(target_date, datetime.max.time(), tzinfo=timezone.utc)

    records = (
        db.query(Transito)
        .filter(Transito.timestamp_in >= start_dt)
        .filter(Transito.timestamp_in <= end_dt)
        .order_by(Transito.timestamp_in.asc())
        .all()
    )

    # Dati completi per un report logistico professionale
    data = []
    for r in records:
        # Calcolo permanenza se uscito
        permanenza = "-"
        if r.timestamp_in and r.timestamp_out:
            diff = r.timestamp_out - r.timestamp_in
            hours, remainder = divmod(diff.total_seconds(), 3600)
            minutes, _ = divmod(remainder, 60)
            permanenza = f"{int(hours)}h {int(minutes)}m"

        data.append({
            "ORA IN": r.timestamp_in.strftime("%H:%M") if r.timestamp_in else "",
            "ORA OUT": r.timestamp_out.strftime("%H:%M") if r.timestamp_out else "-",
            "PERM.": permanenza,
            "TARGA MEZZO": r.targa,
            "TARGA SMR": r.targa_smr if r.targa_smr else "-",
            "VETTORE": r.vettore,
            "OPERAZIONE": r.tipo_operazione.value if r.tipo_operazione else "-",
            "COD. LINEA": r.codice_linea if r.codice_linea else "-",
            "DESTINAZIONE": r.cliente if r.cliente else (r.partenza if r.partenza else "-"),
            "VRID / VRED": r.codice_vred_tme if r.codice_vred_tme else "-",
            "AUTISTA": r.autista if r.autista else "-",
            "TELEFONO": r.telefono if r.telefono else "-",
            "MOLO": r.molo if r.molo else "-",
            "DDT": r.ddt if r.ddt else "-",
            "COLLI": r.colli if r.colli is not None else 0,
            "BI": r.bi if r.bi is not None else "-",
            "CI": r.ci if r.ci is not None else "-",
            "SACCO": r.sacco if r.sacco is not None else "-",
            "SFU %": f"{r.percentuale_sfu}%" if r.percentuale_sfu is not None else "-",
            "PE": r.pe if r.pe is not None else "-",
            "SIGILLI": f"{r.sigillo1 or ''} / {r.sigillo2 or ''}".strip(" /") or "-",
            "STATO": r.stato.value if r.stato else "",
            "NOTE OPERATIVE": r.note if r.note else "-",
        })

    df = pd.DataFrame(data)
    if df.empty:
        df = pd.DataFrame([{"NESSUN DATO": "Nessun transito in questa data"}])

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_name = f"Report_{target_date.strftime('%Y%m%d')}"
        df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=3)
        worksheet = writer.sheets[sheet_name]
        
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import PieChart, Reference
        
        # Colori SDA Corporate
        sda_blue = "0033A0"
        sda_yellow = "FFCC00"
        
        header_fill = PatternFill(start_color=sda_blue, end_color=sda_blue, fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        title_font = Font(color=sda_blue, bold=True, size=18)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        thin_border = Border(left=Side(style='thin', color="D1D5DB"),
                             right=Side(style='thin', color="D1D5DB"),
                             top=Side(style='thin', color="D1D5DB"),
                             bottom=Side(style='thin', color="D1D5DB"))
        
        # Colori Stati Operativi
        colors = {
            "INGRESSO": "3B82F6", # Blu
            "IN_CARICO": "F59E0B", # Ambra
            "COMPLETATO": "10B981", # Verde
            "USCITO": "6B7280"    # Grigio
        }

        # Titolo e Header Grafico
        last_col_letter = get_column_letter(len(df.columns))
        worksheet.merge_cells(f'A1:{last_col_letter}2')
        title_cell = worksheet['A1']
        title_cell.value = f"GATEFLOW - SDA HUB CONTROL | REPORT TRANSITI GIORNALIERO ({target_date.strftime('%d/%m/%Y')})"
        title_cell.font = title_font
        title_cell.alignment = align_center
        
        # Fascia gialla sotto il titolo
        for col in range(1, len(df.columns) + 1):
            worksheet.cell(row=3, column=col).fill = PatternFill(start_color=sda_yellow, end_color=sda_yellow, fill_type="solid")

        # Formatta Header Tabella con Icone Unicode
        header_map = {
            "INGRESSO": "🕒 INGRESSO",
            "USCITA": "📤 USCITA",
            "TARGA": "🚛 TARGA",
            "VETTORE": "🏢 VETTORE",
            "OPERAZIONE": "⚙️ OPERAZIONE",
            "MOLO": "⚓ MOLO",
            "STATO": "📊 STATO",
            "NOTE": "📝 NOTE"
        }
        
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=4, column=col_num)
            raw_val = str(cell.value).upper()
            if raw_val in header_map:
                cell.value = header_map[raw_val]
                
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border

        # Alternanza colori righe
        alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

        # Formatta Righe Tabella
        for r_idx in range(5, len(df) + 5):
            is_alt = r_idx % 2 == 0
            stato_val = ""
            # Trova la colonna STATO per il colore
            stato_col_idx = -1
            for c_idx in range(1, len(df.columns) + 1):
                if "STATO" in str(worksheet.cell(row=4, column=c_idx).value).upper():
                    stato_col_idx = c_idx
                    stato_val = str(worksheet.cell(row=r_idx, column=c_idx).value).strip()
                    break

            for c_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                
                # Colore di sfondo base (alternato)
                if is_alt: cell.fill = alt_fill
                
                # Colorazione cella STATO
                if c_idx == stato_col_idx and stato_val in colors:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill(start_color=colors[stato_val], end_color=colors[stato_val], fill_type="solid")
                
                # Allineamento specifico
                if c_idx in [4, 5, 6, 9, 22, 23]: # Campi testuali lunghi
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

        # Auto-filter e Freeze Panes
        worksheet.auto_filter.ref = f"A4:{last_col_letter}{len(df)+4}"
        worksheet.freeze_panes = "A5"

        # Auto-width delle colonne
        for col_idx, col in enumerate(worksheet.columns, start=1):
            max_len = 0
            column = get_column_letter(col_idx)
            for cell in col:
                if cell.row < 4: continue
                try:
                    if cell.value:
                        length = len(str(cell.value))
                        if length > max_len: max_len = length
                except: pass
            
            # Limiti larghezza
            adjusted_width = min(max(max_len + 5, 12), 45)
            worksheet.column_dimensions[column].width = adjusted_width

        # Footer
        last_row = len(df) + 5
        worksheet.merge_cells(f'A{last_row+1}:{last_col_letter}{last_row+1}')
        footer_cell = worksheet[f'A{last_row+1}']
        footer_text = f"Mezzi totali: {len(df)} | Report generato automaticamente da GateFlow in data {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        footer_cell.value = footer_text
        footer_cell.font = Font(size=10, italic=True, color="6B7280")
        footer_cell.alignment = Alignment(horizontal="right")

        # GRAFICO A TORTA E RIEPILOGO EXECUTIVE
        if not df.columns[0] == "NESSUN DATO" and "STATO" in df.columns:
            stato_counts = df['STATO'].value_counts()
            
            # Tabella di riepilogo visibile
            summary_row = last_row + 3
            summary_col = len(df.columns) - 1 if len(df.columns) > 8 else 8
            
            # Intestazione Riepilogo Executive
            worksheet.cell(row=summary_row-1, column=summary_col, value="EXECUTIVE SUMMARY").font = Font(bold=True, size=12, color=sda_blue)
            
            worksheet.cell(row=summary_row, column=summary_col, value="Stato Operativo")
            worksheet.cell(row=summary_row, column=summary_col+1, value="N° Mezzi")
            
            worksheet.cell(row=summary_row, column=summary_col).font = Font(color="FFFFFF", bold=True)
            worksheet.cell(row=summary_row, column=summary_col).fill = header_fill
            worksheet.cell(row=summary_row, column=summary_col+1).font = Font(color="FFFFFF", bold=True)
            worksheet.cell(row=summary_row, column=summary_col+1).fill = header_fill
            
            for i, (stato, count) in enumerate(stato_counts.items()):
                r = summary_row + 1 + i
                cell_s = worksheet.cell(row=r, column=summary_col, value=stato)
                cell_c = worksheet.cell(row=r, column=summary_col+1, value=count)
                cell_s.border = thin_border
                cell_c.border = thin_border
                cell_c.alignment = align_center
                
                # Colore piccolo indicatore a lato (opzionale, qui coloriamo il testo)
                if stato in colors:
                    cell_s.font = Font(color=colors[stato], bold=True)
                
            pie = PieChart()
            pie.title = "Analisi Distribuzione Hub"
            labels = Reference(worksheet, min_col=summary_col, min_row=summary_row+1, max_row=summary_row+len(stato_counts))
            data_ref = Reference(worksheet, min_col=summary_col+1, min_row=summary_row, max_row=summary_row+len(stato_counts))
            pie.add_data(data_ref, titles_from_data=True)
            pie.set_categories(labels)
            
            # Stile grafico moderno
            pie.style = 10 
            
            # Inserisce il grafico a sinistra del riepilogo
            worksheet.add_chart(pie, f"B{summary_row-1}")

        # IMPOSTAZIONI STAMPA (PRINT READY)
        worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToPage = True
        worksheet.page_setup.fitToHeight = 0 # Adatta solo in larghezza
        worksheet.page_setup.fitToWidth = 1
        
        # Margini stretti per massimizzare lo spazio
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.5
        worksheet.page_margins.bottom = 0.5
        worksheet.page_margins.header = 0.2
        worksheet.page_margins.footer = 0.2

    output.seek(0)
    filename = f"transiti_sda_{target_date.strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
